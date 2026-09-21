import sys, os, re, io
base = "/home/support/INTERCLOUD/backend"
sys.path.insert(0, base)
os.chdir(base)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import portal.routes.business as biz

def render_xlsx(wb):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return biz._preview_xlsx_html(bio.read())

# Test 1: workbook 3 kolom -> only 3 td per row
wb = Workbook()
ws = wb.active
ws["A1"] = "A"
ws["B1"] = "B"
ws["C1"] = "C"
html = render_xlsx(wb)
row = re.search(r"<tr>.*?</tr>", html).group(0)
tds = re.findall(r"<td[^>]*>.*?</td>", row)
assert len(tds) == 3, f"expected 3 td, got {len(tds)} in {row}"
print("PASS: 3-col workbook -> 3 td")

# Test 2: merge A1:B1 -> colspan=2
wb = Workbook()
ws = wb.active
ws.merge_cells("A1:B1")
ws["A1"] = "Merged"
html = render_xlsx(wb)
td = re.search(r"<td[^>]*>Merged</td>", html)
assert td and 'colspan="2"' in td.group(0), f"expected colspan=2, got {td.group(0) if td else None}"
print("PASS: A1:B1 merge -> colspan=2")

# Test 3: merge A1:AZ1 with cap 30 -> colspan=30, not 52
wb = Workbook()
ws = wb.active
ws.merge_cells("A1:AZ1")
ws["A1"] = "Long merged"
html = render_xlsx(wb)
td = re.search(r"<td[^>]*>Long merged</td>", html)
assert td and 'colspan="30"' in td.group(0), f"expected colspan=30, got {td.group(0) if td else None}"
print("PASS: A1:AZ1 merge -> colspan=30")

# Test 4: merge AE1:AF1 (outside 30-col view) -> should NOT create a td at col 30
wb = Workbook()
ws = wb.active
for i in range(1, 31):
    ws.cell(row=1, column=i, value=f"C{i}")
ws.merge_cells("AE1:AF1")  # columns 31-32
ws["AE1"] = "out"
html = render_xlsx(wb)
# row should have exactly 30 td, none containing "out"
row = re.search(r"<tr>.*?</tr>", html).group(0)
tds = re.findall(r"<td[^>]*>.*?</td>", row)
assert len(tds) == 30, f"expected 30 td, got {len(tds)}"
assert "out" not in row, f"out-of-view merge leaked into row: {row}"
print("PASS: AE1:AF1 merge outside view -> not rendered")

# Test 5: merge beyond row 200 -> rowspan clipped to view
wb = Workbook()
ws = wb.active
ws["A1"] = "deep"
ws.merge_cells("A201:A250")  # beyond MAX_ROWS 200
html = render_xlsx(wb)
# A201 is outside view, so A1 row should not have any merged td
row = re.search(r"<tr>.*?</tr>", html).group(0)
tds = re.findall(r"<td[^>]*>.*?</td>", row)
assert len(tds) == 1 and tds[0] == "<td>deep</td>", f"unexpected td for A1 beyond row 200: {row}"
print("PASS: A201:A250 merge outside view -> not rendered")

# Test 6: merge A1:AD1 (30 cols) -> colspan=30 exactly
wb = Workbook()
ws = wb.active
ws.merge_cells("A1:AD1")  # A1:AD1 = 30 columns
ws["A1"] = "full"
html = render_xlsx(wb)
td = re.search(r"<td[^>]*>full</td>", html)
assert td and 'colspan="30"' in td.group(0), f"expected colspan=30 for A1:AD1, got {td.group(0) if td else None}"
print("PASS: A1:AD1 merge -> colspan=30")

print("\nALL XLSX MERGE EDGE CASES PASSED")
