"""
Edge-case tests for _preview_xlsx_html merge/span clipping.

Covers the six assertions Dudung (QA) required before accepting commit 0e9f481:
  1. A1:AZ1 (cols 1-52) → colspan="30" in HTML (clipped to MAX_COLS), not 52.
  2. A1:AZ1 → rowspan attribute absent (merge is row-only).
  3. AE1:AF1 (cols 31-32) → entirely outside view; does NOT appear as colspan
     on column 30.
  4. AE1:AF1 → cell at col 30 with real data still renders in output.
  5. Merge crossing row 200 (e.g. rows 195-205) → rowspan="6" (capped), not 11.
  6. Merge starting at col > MAX_COLS (AG1:AH1) → col-30 data is NOT swallowed.
"""

import io
import re
import sys
import os

import pytest

# Support both containerised (/app/backend) and local dev paths.
for _p in ("/app/backend", "/home/support/INTERCLOUD/backend"):
    if _p not in sys.path:
        sys.path.insert(0, _p)
os.chdir(os.path.dirname(os.path.abspath(__file__)) + "/..")

import portal.routes.business as biz  # noqa: E402


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _xlsx_bytes(setup_fn) -> bytes:
    """Create an in-memory xlsx with setup_fn(ws) applied to the active sheet."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    setup_fn(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _html(setup_fn) -> str:
    return biz._preview_xlsx_html(_xlsx_bytes(setup_fn))


# ---------------------------------------------------------------------------
# test 1 & 2 — A1:AZ1 (cols 1–52) clipped to colspan=30, no rowspan
# ---------------------------------------------------------------------------

def _wide_merge(ws):
    ws["A1"] = "wide header"
    ws.merge_cells("A1:AZ1")  # col 1 → col 52; MAX_COLS = 30


def test_wide_merge_colspan_capped_at_30():
    """A1:AZ1 → colspan must be 30, not 52 (assertion 1)."""
    html = _html(_wide_merge)
    # Find all colspan values in the output
    colspans = [int(v) for v in re.findall(r'colspan="(\d+)"', html)]
    assert colspans, "Expected at least one colspan in output"
    assert max(colspans) == 30, (
        f"colspan should be 30 (MAX_COLS), got {max(colspans)}"
    )


def test_wide_merge_no_rowspan():
    """A1:AZ1 is a single-row merge → no rowspan attribute (assertion 2)."""
    html = _html(_wide_merge)
    assert 'rowspan=' not in html, (
        "Single-row horizontal merge must not emit rowspan"
    )


# ---------------------------------------------------------------------------
# test 3 & 4 — AE1:AF1 (cols 31–32) fully outside view
# ---------------------------------------------------------------------------

def _out_of_view_col_merge(ws):
    ws.cell(row=1, column=30).value = "visible-col30"   # col 30 — IN view
    ws.merge_cells("AE1:AF1")                            # cols 31–32 — OUT of view
    ws["AE1"] = "outside"


def test_out_of_view_merge_not_mapped_to_col30():
    """AE1:AF1 must not produce colspan on col 30 (assertion 3)."""
    html = _html(_out_of_view_col_merge)
    # If the bug were present the out-of-view merge's origin would be clamped
    # to col 30, turning the col-30 cell into a phantom master with colspan.
    # After the fix there should be no colspan at all (the merge is skipped).
    assert 'colspan=' not in html, (
        "Out-of-view merge AE1:AF1 must not produce any colspan in the HTML"
    )


def test_col30_data_survives_out_of_view_merge():
    """Data at col 30 must still render even when a merge starts at col >30 (assertion 4)."""
    html = _html(_out_of_view_col_merge)
    assert "visible-col30" in html, (
        "Cell at col 30 should be visible; it must not be swallowed by an "
        "out-of-view merge"
    )


# ---------------------------------------------------------------------------
# test 5 — merge crossing row 200 → rowspan capped to view boundary
# ---------------------------------------------------------------------------

def _row_boundary_merge(ws):
    # Put data in rows 195–205; merge rows 195–205 in col 1
    for r in range(195, 206):
        ws.cell(row=r, column=1).value = f"r{r}"
    ws.merge_cells(start_row=195, start_column=1, end_row=205, end_column=1)
    ws["A195"] = "cross-boundary"


def test_rowspan_capped_at_view_boundary():
    """Merge rows 195-205 with MAX_ROWS=200 → rowspan=6, not 11 (assertion 5)."""
    html = _html(_row_boundary_merge)
    rowspans = [int(v) for v in re.findall(r'rowspan="(\d+)"', html)]
    assert rowspans, "Expected a rowspan in output for a multi-row merge"
    assert max(rowspans) == 6, (
        f"rowspan should be 6 (rows 195-200 = 6 rows), got {max(rowspans)}"
    )


# ---------------------------------------------------------------------------
# test 6 — merge starting at col > MAX_COLS must not hide col-30 data
# ---------------------------------------------------------------------------

def _merge_starts_beyond_cap(ws):
    ws.cell(row=1, column=30).value = "col30-safe"
    ws.merge_cells("AG1:AH1")   # cols 33–34; origin > 30 → skip entirely
    ws["AG1"] = "beyond cap"


def test_col30_safe_when_merge_origin_beyond_cap():
    """Merge starting at col > MAX_COLS must not shadow col-30 data (assertion 6)."""
    html = _html(_merge_starts_beyond_cap)
    assert "col30-safe" in html, (
        "Col-30 value must appear; a merge whose origin is beyond MAX_COLS "
        "must be silently dropped, not clamped into the view"
    )
